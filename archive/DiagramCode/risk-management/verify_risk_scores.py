# -*- coding: utf-8 -*-
"""
Независимая проверка итогового коэффициента риска (вероятность x воздействие)
для реестра рисков DiagramCode — без доверия формулам Excel, чистым Python.
Читает сырые значения F/G прямо из файла (а не из захардкоженной копии), затем
пересчитывает H = F*G сама, поэтому при правках "Карта рисков - DiagramCode.xlsx"
скрипт всегда сверяется с текущим состоянием файла, а не с его снимком на момент
написания.

Запуск: python verify_risk_scores.py
"""

import os
import openpyxl

XLSX_PATH = os.path.join(os.path.dirname(__file__), "Карта рисков - DiagramCode.xlsx")
SHEET = "Таблица с рисками"
FIRST_ROW = 4
LAST_ROW = 23  # конец диапазона таблицы Таблица1; пустые строки просто пропускаются


def load_risks(path=XLSX_PATH):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[SHEET]
    risks = []
    for row in range(FIRST_ROW, LAST_ROW + 1):
        name = ws[f"C{row}"].value
        f = ws[f"F{row}"].value
        g = ws[f"G{row}"].value
        if name is None or f is None or g is None:
            continue
        risks.append((name, float(f), float(g)))
    return risks


if __name__ == "__main__":
    risks = load_risks()
    ranked = sorted(risks, key=lambda r: r[1] * r[2], reverse=True)
    print(f"Прочитано рисков из файла: {len(risks)}")
    print(f"{'#':>2} | {'F':>4} | {'G':>4} | {'H=F*G':>6} | Риск")
    print("-" * 90)
    for i, (name, f, g) in enumerate(ranked, 1):
        print(f"{i:>2} | {f:>4} | {g:>4} | {f * g:>6.2f} | {name}")
