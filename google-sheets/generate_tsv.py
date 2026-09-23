# -*- coding: utf-8 -*-
"""
Генерирует готовые к импорту в Google Таблицы файлы (.tsv) из тех же данных,
что используются в проверочных скриптах verify_cpm*.py и в Excel-файле рисков.
Один источник данных — поэтому в Google Таблицах получатся ровно те же числа,
что уже проверены в Excel и Python.

Запуск: python generate_tsv.py
"""

import importlib.util
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)


def load_module(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def col(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def build_cpm_sheet(tasks):
    order = sorted(tasks)
    row_of = {t: i + 2 for i, t in enumerate(order)}
    first, last = 2, len(order) + 1
    succ = {t: [] for t in tasks}
    for t, info in tasks.items():
        for p in info["pred"]:
            succ[p].append(t)

    duration = 0
    ef = {}
    for t in order:
        es = 1 if not tasks[t]["pred"] else max(ef[p] for p in tasks[t]["pred"])
        ef[t] = es + tasks[t]["dur"]
    duration = max(ef.values()) - 1

    gantt_first_col = 12  # L
    days = list(range(1, duration + 1))

    header = ["№", "Задача", "Предшественники", "Длительность", "ES", "EF", "LS", "LF",
              "Резерв", "Критический путь", ""]
    header += [str(d) for d in days]
    lines = ["\t".join(header)]

    for t in order:
        r = row_of[t]
        preds = tasks[t]["pred"]
        if not preds:
            es = "1"
        elif len(preds) == 1:
            es = f"=F{row_of[preds[0]]}"
        else:
            es = "=MAX(" + ",".join(f"F{row_of[p]}" for p in preds) + ")"
        if not succ[t]:
            lf = f"=MAX($F${first}:$F${last})"
        elif len(succ[t]) == 1:
            lf = f"=G{row_of[succ[t][0]]}"
        else:
            lf = "=MIN(" + ",".join(f"G{row_of[s]}" for s in succ[t]) + ")"
        cells = [
            str(t),
            tasks[t]["name"],
            ", ".join(str(p) for p in preds) or "—",
            str(tasks[t]["dur"]),
            es,
            f"=E{r}+D{r}",
            f"=H{r}-D{r}",
            lf,
            f"=G{r}-E{r}",
            f'=IF(I{r}=0,"ДА","НЕТ")',
            "",
        ]
        for i in range(len(days)):
            c = col(gantt_first_col + i)
            cells.append(f'=IF(AND({c}$1>=$E{r},{c}$1<$F{r}),IF($J{r}="ДА","█","░"),"")')
        lines.append("\t".join(cells))

    lines.append("")
    lines.append(f"Общая длительность проекта, рабочих дней\t=MAX(F{first}:F{last})-MIN(E{first}:E{last})")
    lines.append(f'Критический путь\t=TEXTJOIN(" → ",TRUE,FILTER(A{first}:A{last},J{first}:J{last}="ДА"))')
    lines.append(f'Проверка: сумма длительностей работ критического пути\t=SUMIF(J{first}:J{last},"ДА",D{first}:D{last})')
    lines.append("Легенда графика Ганта (столбцы L и дальше)\t█ — работа на критическом пути, ░ — работа с резервом")
    return "\n".join(lines) + "\n", duration


def build_risk_sheet():
    path = os.path.join(ROOT, "risk-management", "Карта рисков - DiagramCode.xlsx")
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["Таблица с рисками"]
    header = ["№", "Обозначение", "Риск", "Влияние на проект", "Триггер",
              "Вероятность (F)", "Воздействие (G)", "Коэффициент H = F×G", "Приоритет",
              "Место по H", "Меры по минимизации"]
    lines = ["\t".join(header)]
    rows = []
    for r in range(4, 24):
        name = ws[f"C{r}"].value
        if name is None:
            continue
        rows.append((ws[f"A{r}"].value, name, ws[f"D{r}"].value, ws[f"E{r}"].value,
                     ws[f"F{r}"].value, ws[f"G{r}"].value, ws[f"I{r}"].value))
    first, last = 2, len(rows) + 1
    for i, (num, name, impact, trigger, f, g, measures) in enumerate(rows):
        r = i + 2
        cells = [
            str(num),
            f'="R"&A{r}',
            name,
            impact,
            trigger,
            str(f),
            str(g),
            f"=F{r}*G{r}",
            f'=IF(H{r}>=0.505,"Высокий",IF(H{r}>=0.14,"Средний","Низкий"))',
            f"=RANK(H{r},$H${first}:$H${last})",
            measures,
        ]
        for c in cells:
            assert "\t" not in c and "\n" not in c
        lines.append("\t".join(cells))
    lines.append("")
    lines.append("Пороги приоритета взяты из легенды шаблона «Карта рисков»: высокий — от 0,505, средний — от 0,14 до 0,505, низкий — меньше 0,14")
    return "\n".join(lines) + "\n", len(rows)


def main():
    gantt_dir = os.path.join(ROOT, "pz-wbs-gantt")
    a = load_module(os.path.join(gantt_dir, "verify_cpm.py"), "verify_cpm")
    b = load_module(os.path.join(gantt_dir, "verify_cpm_example1.py"), "verify_cpm_example1")

    outputs = []
    text, dur = build_cpm_sheet(a.tasks)
    outputs.append(("1 - WBS-Гант - DiagramCode.tsv", text, f"12 работ, {dur} рабочих дней"))
    text, dur = build_cpm_sheet(b.tasks)
    outputs.append(("2 - WBS-Гант - Пример 1 (методичка).tsv", text, f"12 работ, {dur} рабочих дней"))
    text, n = build_risk_sheet()
    outputs.append(("3 - Риски - DiagramCode.tsv", text, f"{n} рисков"))

    for fname, text, info in outputs:
        with open(os.path.join(HERE, fname), "w", encoding="utf-8", newline="\n") as fh:
            fh.write(text)
        print(f"{fname}: {info}")


if __name__ == "__main__":
    main()
