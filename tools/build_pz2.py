import os
import subprocess
import sys

import matplotlib

matplotlib.use('Agg')
import matplotlib.pyplot as plt
import openpyxl
from matplotlib.patches import Patch, Rectangle

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_base import ROOT, Report

XLSX_TEMPLATE = os.path.join(ROOT, 'ПЗ2_WBS-и-диаграмма-Ганта', '02_задание-оригиналы', 'Планировщик проекта на основе диаграммы Ганта - 60121.xlsx')
OUT_DIR = os.path.join(ROOT, 'ПЗ2_WBS-и-диаграмма-Ганта', '07_результаты', 'пример-1-методички')
XLSX_OUT = os.path.join(OUT_DIR, 'Планировщик проекта - Пример 1 - Хаджинова.xlsx')
REPORT_OUT = os.path.join(OUT_DIR, 'Отчёт - ПЗ2 - WBS и график Ганта - Хаджинова.docx')
WBS_DOT = os.path.join(OUT_DIR, 'WBS - Пример 1.dot')
WBS_PNG = os.path.join(OUT_DIR, 'WBS - Пример 1.png')
NET_DOT = os.path.join(OUT_DIR, 'Сетевой график - Пример 1.dot')
NET_PNG = os.path.join(OUT_DIR, 'Сетевой график - Пример 1.png')
GANTT_PNG = os.path.join(OUT_DIR, 'Планировщик проекта - диаграмма Ганта.png')

PROJECT = 'Обучение сотрудников кибербезопасности'
FIRST_ROW = 8
PERIOD = 60

STAGES = [
    ('1', 'Подготовка', [1, 2, 3]),
    ('2', 'Настройка платформы', [4, 5, 6]),
    ('3', 'Разработка контента и обучение', [7, 8, 9, 10]),
    ('4', 'Контроль и внедрение', [11, 12]),
]

TASKS = [
    (1, 'Устав и стартовое совещание', [], 3, 'Утверждение устава, совещание с руководителями'),
    (2, 'Анализ регламентов ИБ', [1], 5, 'Изучение внутренних документов по защите ПДн'),
    (3, 'Фишинг-замер + отчёт', [2], 4, 'Скрытая рассылка, обработка результатов'),
    (4, 'Выбор вендора + договор', [1], 6, 'Анализ рынка, переговоры, подписание'),
    (5, 'Учётные записи (200 чел.)', [4], 3, 'Импорт сотрудников в платформу'),
    (6, 'Настройка white-list', [3, 4], 3, 'Внесение IP-адресов в безопасный список почты'),
    (7, 'Курс по ПДн (№ 99-З)', [6], 8, 'Разработка юридического модуля обучения'),
    (8, 'Модули, видео, тесты', [6], 10, 'Создание интерактивного контента по фишингу'),
    (9, 'Информирование сотрудников', [7, 8], 2, 'Письмо директора, инструкции'),
    (10, 'Микрообучение (4 недели)', [9], 20, 'Основная фаза обучения (4 недели × 5 дней)'),
    (11, 'Итоговый фишинг + тест', [10], 4, 'Контрольный замер, проверка знаний'),
    (12, 'Анализ KPI + отчёт + регламент', [11], 5, 'Дашборд KPI, утверждение регламента'),
]

FULL_NAMES = {
    1: 'Разработка и утверждение устава проекта, проведение стартового совещания',
    2: 'Анализ существующих регламентов информационной безопасности',
    3: 'Проведение скрытой фишинг-атаки (замер «нулевого» уровня) и формирование отчёта об уязвимостях',
    4: 'Анализ рынка, выбор вендора и заключение договора с оплатой лицензии',
    5: 'Создание учётных записей для 200 сотрудников',
    6: 'Настройка интеграции с почтовым сервером (white-list IP-адресов)',
    7: 'Подготовка курса по Закону РБ № 99-З «О защите персональных данных»',
    8: 'Разработка интерактивных модулей по фишингу и социальной инженерии, видеолекций и базы тестов',
    9: 'Информирование сотрудников о старте программы',
    10: 'Проведение активной фазы микрообучения (4 недели) и семинаров для отстающих',
    11: 'Проведение итоговой контрольной фишинг-атаки и финального тестирования',
    12: 'Анализ достижения KPI, подготовка отчёта для директора, утверждение регламента регулярного обучения',
}


def cpm():
    by_id = {t[0]: t for t in TASKS}
    succ = {t[0]: [s[0] for s in TASKS if t[0] in s[2]] for t in TASKS}
    es, ef = {}, {}
    for tid, _, preds, dur, _ in TASKS:
        es[tid] = max((ef[p] for p in preds), default=1)
        ef[tid] = es[tid] + dur
    finish = max(ef.values())
    ls, lf = {}, {}
    for tid, _, _, dur, _ in reversed(TASKS):
        lf[tid] = min((ls[s] for s in succ[tid]), default=finish)
        ls[tid] = lf[tid] - dur
    tf = {t: ls[t] - es[t] for t in es}
    return by_id, succ, es, ef, ls, lf, tf, finish


BY_ID, SUCC, ES, EF, LS, LF, TF, FINISH = cpm()
CRITICAL = [t for t in ES if TF[t] == 0]
DURATION = FINISH - 1


def fmt_ids(ids):
    return ', '.join(str(i) for i in ids) if ids else '—'


def row_of(tid):
    return FIRST_ROW + tid - 1


def build_xlsx():
    wb = openpyxl.load_workbook(XLSX_TEMPLATE)
    ws = wb['ТАБЛИЦА ДАННЫХ']
    last_row = row_of(len(TASKS))
    for tid, name, preds, dur, _ in TASKS:
        r = row_of(tid)
        ws[f'B{r}'] = tid
        ws[f'C{r}'] = name
        ws[f'D{r}'] = fmt_ids(preds)
        ws[f'E{r}'] = fmt_ids(SUCC[tid])
        if not preds:
            ws[f'F{r}'] = 1
        elif len(preds) == 1:
            ws[f'F{r}'] = f'=H{row_of(preds[0])}'
        else:
            ws[f'F{r}'] = '=MAX(' + ','.join(f'H{row_of(p)}' for p in preds) + ')'
        ws[f'G{r}'] = dur
        ws[f'H{r}'] = f'=F{r}+G{r}'
        succ = SUCC[tid]
        if not succ:
            ws[f'J{r}'] = f'=MAX($H${FIRST_ROW}:$H${last_row})'
        elif len(succ) == 1:
            ws[f'J{r}'] = f'=I{row_of(succ[0])}'
        else:
            ws[f'J{r}'] = '=MIN(' + ','.join(f'I{row_of(s)}' for s in succ) + ')'
    for r in range(last_row + 1, 23):
        for col in 'BCDEFGHIJKL':
            ws[f'{col}{r}'] = None
    wb.save(XLSX_OUT)


def run_dot(src, dst, text):
    with open(src, 'w', encoding='utf-8') as f:
        f.write(text)
    subprocess.run(['dot', '-Tpng', '-Gdpi=160', src, '-o', dst], check=True)


def build_wbs():
    lines = [
        'digraph WBS {',
        '  graph [rankdir=TB, splines=ortho, nodesep=0.3, ranksep=0.45];',
        '  node [shape=box, style="rounded,filled", fontname="Arial", fontsize=11];',
        '  edge [arrowhead=none];',
        f'  root [label="0 {PROJECT}\\n(Пример 1)", fillcolor="#d9e2f3", fontsize=13];',
    ]
    for code, name, ids in STAGES:
        lines.append(f'  s{code} [label="{code}.0\\n{name}", fillcolor="#e2efd9"];')
        lines.append(f'  root -> s{code};')
        prev = f's{code}'
        for k, tid in enumerate(ids, 1):
            label = BY_ID[tid][1].replace(' + ', '\\n+ ')
            lines.append(f'  t{tid} [label="{code}.{k} (работа {tid})\\n{label}", fillcolor="#fff2cc", fontsize=10];')
            lines.append(f'  {prev} -> t{tid};')
            prev = f't{tid}'
    lines.append('}')
    run_dot(WBS_DOT, WBS_PNG, '\n'.join(lines))


def build_network():
    lines = [
        'digraph NET {',
        '  graph [rankdir=LR, nodesep=0.35, ranksep=0.55];',
        '  node [shape=plain, fontname="Arial", fontsize=10];',
        '  edge [color="#555555"];',
    ]
    for tid, name, preds, dur, _ in TASKS:
        crit = TF[tid] == 0
        color = '#f4b183' if crit else '#bdd7ee'
        border = '#c00000' if crit else '#2f5597'
        cut = name.rfind(' ', 0, 20)
        short = name if len(name) <= 20 or cut < 0 else name[:cut] + '<br/>' + name[cut + 1:]
        label = (
            f'<<table border="1" cellborder="1" cellspacing="0" cellpadding="3" color="{border}" bgcolor="white">'
            f'<tr><td>{ES[tid]}</td><td bgcolor="{color}"><b>{tid}</b></td><td>{EF[tid]}</td></tr>'
            f'<tr><td colspan="3">{short}</td></tr>'
            f'<tr><td>{LS[tid]}</td><td>D={dur}, R={TF[tid]}</td><td>{LF[tid]}</td></tr></table>>'
        )
        lines.append(f'  n{tid} [label={label}];')
        for p in preds:
            style = 'color="#c00000", penwidth=2.2' if crit and TF[p] == 0 else ''
            lines.append(f'  n{p} -> n{tid} [{style}];')
    lines.append('}')
    run_dot(NET_DOT, NET_PNG, '\n'.join(lines))


def build_gantt():
    n = len(TASKS)
    fig, ax = plt.subplots(figsize=(16, 6.2), dpi=150)
    for day in range(1, PERIOD + 1):
        shade = '#f2f2f2' if day % 2 else '#ffffff'
        ax.add_patch(Rectangle((day - 1, 0), 1, n, facecolor=shade, edgecolor='#d9d9d9', linewidth=0.4))
    for i, (tid, name, _, dur, _) in enumerate(TASKS):
        y = n - i - 1
        color = '#ed7d31' if TF[tid] == 0 else '#5b9bd5'
        for day in range(ES[tid], ES[tid] + dur):
            ax.add_patch(Rectangle((day - 1, y + 0.15), 1, 0.7, facecolor=color, edgecolor='white', linewidth=0.3))
    ax.set_xlim(0, PERIOD)
    ax.set_ylim(0, n)
    ax.set_yticks([n - i - 0.5 for i in range(n)])
    ax.set_yticklabels([f'{t[0]}. {t[1]}' for t in TASKS], fontsize=10)
    ax.set_xticks([d - 0.5 for d in range(1, PERIOD + 1)])
    ax.set_xticklabels([str(d) for d in range(1, PERIOD + 1)], fontsize=7)
    ax.xaxis.tick_top()
    ax.tick_params(length=0)
    for side in ax.spines.values():
        side.set_visible(False)
    ax.set_xlabel('Периоды (рабочие дни), J4:BQ4', fontsize=10)
    ax.legend(handles=[Patch(color='#ed7d31', label='Критический путь (I = 100 %)'),
                       Patch(color='#5b9bd5', label='Работа с резервом (I = 0)')],
              loc='lower center', bbox_to_anchor=(0.5, -0.16), ncol=2, fontsize=10, frameon=False)
    fig.tight_layout()
    fig.savefig(GANTT_PNG)
    plt.close(fig)


def build_report():
    rep = Report(2, 'Построение WBS и графика Ганта')
    rep.heading1('1 Цель работы', new_page=False)
    rep.text(
        'Освоение методов календарно-сетевого планирования: построение иерархической структуры работ (WBS), '
        'определение длительностей и взаимосвязей работ, расчёт ранних и поздних сроков методом критического '
        'пути (CPM) и построение диаграммы Ганта в Excel-планировщике на данных примера 1 методических указаний '
        f'(проект «{PROJECT}»).'
    )

    rep.heading1('2 Ход работы')
    rep.heading2('2.1 Теоретические сведения')
    rep.text(
        'Иерархическая структура работ (Work Breakdown Structure, WBS) — иерархическая декомпозиция проекта '
        'на подпроекты, пакеты работ и отдельные работы. Работы, не включённые в WBS, не являются работами проекта.'
    )
    rep.text(
        'Сетевой график строится методом предшествования (Precedence Diagramming Method, PDM): работы '
        'изображаются узлами, зависимости «финиш – старт» — стрелками. Прямой проход определяет ранний старт '
        'ES = max(EF предшественников) и ранний финиш EF = ES + D; обратный проход — поздний финиш '
        'LF = min(LS последователей) и поздний старт LS = LF – D. Полный резерв TF = LS – ES. Работы '
        'с нулевым резервом образуют критический путь, длительность которого равна длительности проекта.'
    )
    rep.text(
        'Диаграмма Ганта состоит из горизонтальных полос вдоль оси времени: каждая полоса соответствует '
        'работе, её концы — моментам начала и завершения, протяжённость — длительности.'
    )

    rep.heading2('2.2 Исходные данные')
    rep.text(
        f'Исходные данные взяты с листа «Пример 1» файла «Планировщик проекта на основе диаграммы Ганта – 60121.xlsx»: '
        f'проект «{PROJECT}», 4 этапа, 12 работ. Состав этапов приведён в таблице {rep.next_table()}.'
    )
    rows = []
    for code, name, ids in STAGES:
        for tid in ids:
            rows.append([f'{code}. {name}', tid, FULL_NAMES[tid]])
    rep.table('Этапы и работы проекта', ['Этап', '№', 'Содержание работы'], rows, [4.5, 1.2, 11])

    rep.heading2('2.3 Иерархическая структура работ')
    rep.text(
        f'WBS построена по подходу жизненного цикла (декомпозиция по этапам создания результата). Уровень 0 — '
        f'проект, уровень 1 — четыре этапа с кодами 1.0–4.0, уровень 2 — двенадцать работ (рисунок {rep.next_figure()}).'
    )
    rep.figure(WBS_PNG, f'Иерархическая структура работ проекта «{PROJECT}»')

    rep.heading2('2.4 Длительности и взаимосвязи работ')
    rep.text(
        f'Длительности заданы в рабочих днях по данным примера 1; все зависимости имеют тип «финиш – старт». '
        f'Непосредственные предшественники и последователи приведены в таблице {rep.next_table()}.'
    )
    rep.table(
        'Длительности и взаимосвязи работ',
        ['№', 'Работа', 'Предше-ственники', 'Последо-ватели', 'D, дн.', 'Обоснование длительности'],
        [[t[0], t[1], fmt_ids(t[2]), fmt_ids(SUCC[t[0]]), t[3], t[4]] for t in TASKS],
        [1, 4.6, 2.2, 2.2, 1.4, 6],
    )
    rep.text(
        'Работа 6 является операцией слияния (два предшественника — 3 и 4), работа 9 — также операцией '
        'слияния (7 и 8). Работы 1, 4 и 6 являются дробящимися операциями. Работы 7 и 8 выполняются параллельно.'
    )

    rep.heading2('2.5 Расчёт методом критического пути')
    rep.text(
        f'Результаты прямого и обратного проходов приведены в таблице {rep.next_table()}. Отсчёт ведётся '
        'с ES = 1, поэтому EF обозначает первый день после завершения работы.'
    )
    rep.table(
        'Расчёт параметров сетевого графика',
        ['№', 'Работа', 'D', 'ES', 'EF', 'LS', 'LF', 'TF', 'Крит.'],
        [[t[0], t[1], t[3], ES[t[0]], EF[t[0]], LS[t[0]], LF[t[0]], TF[t[0]], 'ДА' if TF[t[0]] == 0 else 'НЕТ']
         for t in TASKS],
        [1, 5.8, 1.1, 1.2, 1.2, 1.2, 1.2, 1.2, 1.6],
    )
    crit_path = ' → '.join(str(t) for t in CRITICAL)
    crit_sum = ' + '.join(str(BY_ID[t][3]) for t in CRITICAL)
    reserves = ', '.join(f'{t} ({TF[t]} дн.)' for t in ES if TF[t] > 0)
    rep.text(
        f'Критический путь: {crit_path}. Длительность проекта: EF(12) – ES(1) = {FINISH} – 1 = {DURATION} '
        f'рабочих дней; контроль по сумме длительностей работ критического пути: {crit_sum} = {DURATION}. '
        f'Работы с резервом: {reserves}. Работа 5 не имеет последователей, поэтому её поздний финиш '
        f'равен сроку завершения проекта, а резерв максимален.'
    )
    rep.text(f'Сетевой график с параметрами работ приведён на рисунке {rep.next_figure()}.')
    rep.figure(NET_PNG, 'Сетевой график проекта (сверху: ES, №, EF; снизу: LS, длительность и резерв, LF); '
                        'критический путь выделен красным')

    rep.heading2('2.6 Заполнение листа «ТАБЛИЦА ДАННЫХ»')
    rep.text(
        'Работы внесены в строки 8–19 структурированной таблицы «Таблица1». Содержимое столбцов:'
    )
    rep.bullets([
        'B — номер работы; C — наименование; D, E — предшественники и последователи (справочно);',
        'F (ES): 1 для работы без предшественников, =H{R} для одного предшественника, =MAX(H{R1},H{R2}) для нескольких;',
        'G (D) — длительность; H (EF): =F+G;',
        'J (LF): =MAX($H$8:$H$19) для завершающих работ (5 и 12), =I{R} для одного последователя, '
        '=MIN(I{R1},I{R2}) для нескольких;',
        'I (LS): =J–G; K (TF): =I–F; L: =IF(K=0,"ДА","НЕТ ").',
    ])
    rep.text('Строки 20–22 шаблона очищены; листы «Пример 1» и «Пример 2» не изменены.')

    rep.heading2('2.7 Проверка листа «Планировщик проекта»')
    rep.text(
        'Лист «Планировщик проекта» получает данные формулами со смещением на 3 строки: столбец C — наименование '
        '(C!C8), F — резерв (K8), G — раннее начало (F8), H — длительность (G8), I — признак критического пути '
        '(100 % при L8 = «ДА»). Правила условного форматирования диапазона J5:BQ30 закрашивают ячейку периода '
        'J$4, если J$4 ∈ [G; G + H – 1]; для работ критического пути применяется отдельный цвет. '
        f'Рассматриваемый период J2 = {PERIOD} дней вмещает весь проект ({DURATION} дней).'
    )
    rep.text(
        f'Диаграмма, формируемая этими правилами по рассчитанным значениям, приведена на рисунке {rep.next_figure()}: '
        '12 полос, положение и длина каждой полосы совпадают со значениями ES и D таблицы 3, полосы критического '
        'пути выделены цветом.'
    )
    rep.figure(GANTT_PNG, 'Диаграмма Ганта листа «Планировщик проекта» (правила условного форматирования J5:BQ30)')

    rep.heading1('Выводы')
    rep.text(
        f'Для проекта «{PROJECT}» построена иерархическая структура работ из 4 этапов и 12 работ, определены '
        'длительности и зависимости «финиш – старт», методом критического пути рассчитаны ранние и поздние '
        f'сроки и резервы. Критический путь {crit_path} определяет длительность проекта {DURATION} рабочих дней; '
        f'работы 4, 5 и 7 имеют резерв. Лист «ТАБЛИЦА ДАННЫХ» Excel-планировщика заполнен формулами CPM, '
        'на листе «Планировщик проекта» сформирована диаграмма Ганта с выделением критического пути.'
    )
    rep.save(REPORT_OUT)


def verify():
    expected = {1: (1, 4, 1, 4, 0), 5: (10, 13, 54, 57, 44), 10: (28, 48, 28, 48, 0), 12: (52, 57, 52, 57, 0)}
    for tid, vals in expected.items():
        assert (ES[tid], EF[tid], LS[tid], LF[tid], TF[tid]) == vals, tid
    assert CRITICAL == [1, 2, 3, 6, 8, 9, 10, 11, 12]
    assert DURATION == 56


if __name__ == '__main__':
    verify()
    os.makedirs(OUT_DIR, exist_ok=True)
    build_xlsx()
    build_wbs()
    build_network()
    build_gantt()
    build_report()
    print('ok', DURATION, CRITICAL)
