"""Отчёт по практическому занятию в оформлении СТП 01–2017 БГУИР.

Правила (СТП, раздел 2 и приложения Л, Н):
  - А4, поля: левое 30 мм, правое 15 мм, верхнее и нижнее 20 мм; номер страницы — внизу справа, 10 мм от края;
    на титульном листе номер не ставится;
  - Times New Roman 14 пт, межстрочный интервал 1,0, абзацный отступ 1,25 см, выравнивание по ширине;
  - заголовки разделов — прописными, полужирным, с абзацного отступа, без точки; раздел — с новой страницы;
    заголовки подразделов — строчными с прописной, полужирным; между заголовком и текстом — пробельная строка;
  - ненумерованные разделы (СОДЕРЖАНИЕ, ВЫВОДЫ, СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ, ПРИЛОЖЕНИЕ) — по центру;
  - рисунок — по центру, пробельная строка до и после, подпись «Рисунок 2.1 – Наименование» по центру;
  - таблица — «Таблица 2.1 – Заголовок» над таблицей от левой границы, таблица отделена пробельными строками;
  - формулы — по центру, номер (2.1) у правого края;
  - перечисления — с абзацного отступа со знаком «–», в конце «;», последний элемент — «.»;
  - ссылки на источники — [n], список источников по ГОСТ 7.1, приложения — «ПРИЛОЖЕНИЕ А (обязательное)».
Содержание формируется полем TOC; номера страниц в нём проставляет Word (скрипт word_finalize.ps1).
"""
import os

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = 'Times New Roman'
TEXT_WIDTH_CM = 21.0 - 3.0 - 1.5


def _font(run, size=14, bold=False, italic=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(attr), name)
    return run


def _field(paragraph, instr, placeholder=''):
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar')
    b.set(qn('w:fldCharType'), 'begin')
    run._r.append(b)
    run2 = paragraph.add_run()
    t = OxmlElement('w:instrText')
    t.set(qn('xml:space'), 'preserve')
    t.text = instr
    run2._r.append(t)
    run3 = paragraph.add_run()
    s = OxmlElement('w:fldChar')
    s.set(qn('w:fldCharType'), 'separate')
    run3._r.append(s)
    run4 = paragraph.add_run(placeholder)
    run5 = paragraph.add_run()
    e = OxmlElement('w:fldChar')
    e.set(qn('w:fldCharType'), 'end')
    run5._r.append(e)
    for r in (run, run2, run3, run4, run5):
        _font(r)


def _borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tbl_pr.append(borders)


def _cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    tr_pr.append(el)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader')
    h.set(qn('w:val'), 'true')
    tr_pr.append(h)


class STPReport:
    def __init__(self):
        self.doc = docx.Document()
        self._setup_styles()
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.footer_distance = Cm(1.0)
        sec.header_distance = Cm(1.0)
        sec.different_first_page_header_footer = True
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.first_line_indent = Cm(0)
        _field(fp, 'PAGE', '1')
        self.section_no = 0
        self.letter = None
        self.fig = 0
        self.tab = 0
        self.frm = 0
        self.sources = []

    def _setup_styles(self):
        st = self.doc.styles['Normal']
        st.font.name = FONT
        st.font.size = Pt(14)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.first_line_indent = Cm(1.25)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True
        for name, size in (('Heading 1', 14), ('Heading 2', 14)):
            h = self.doc.styles[name]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = None
            rpr = h.element.get_or_add_rPr()
            fonts = rpr.find(qn('w:rFonts'))
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                rpr.append(fonts)
            for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
                fonts.set(qn(attr), FONT)
            for tag in ('w:color', 'w:sz', 'w:szCs'):
                for el in rpr.findall(qn(tag)):
                    if tag == 'w:color':
                        rpr.remove(el)
            h.paragraph_format.space_before = Pt(0)
            h.paragraph_format.space_after = Pt(0)
            h.paragraph_format.line_spacing = 1.0
            h.paragraph_format.first_line_indent = Cm(1.25)
            h.paragraph_format.left_indent = Cm(0)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # --- базовые элементы -------------------------------------------------
    def p(self, text='', align=None, indent=True, size=14, bold=False, italic=False, keep=False):
        par = self.doc.add_paragraph()
        if align is not None:
            par.alignment = align
        if not indent:
            par.paragraph_format.first_line_indent = Cm(0)
        if keep:
            par.paragraph_format.keep_with_next = True
        if text:
            _font(par.add_run(text), size, bold, italic)
        return par

    def blank(self, keep=False):
        return self.p('', keep=keep)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- титульный лист ---------------------------------------------------
    def title_page(self, number, topic, student='К. А. Хаджинова', group='60121', teacher='И. В. Кашникова', year=2026):
        c = WD_ALIGN_PARAGRAPH.CENTER
        for line in ('Министерство образования Республики Беларусь',
                     'Учреждение образования',
                     'БЕЛОРУССКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ',
                     'ИНФОРМАТИКИ И РАДИОЭЛЕКТРОНИКИ',
                     'Обособленное подразделение «Институт информационных технологий БГУИР»'):
            self.p(line, c, indent=False)
        self.blank()
        self.p('Факультет повышения квалификации и переподготовки', c, indent=False)
        self.p('Кафедра микропроцессорных систем и сетей', c, indent=False)
        for _ in range(7):
            self.blank()
        self.p('ОТЧЕТ', c, indent=False, bold=True)
        self.p(f'по практическому занятию № {number}', c, indent=False)
        self.p('по дисциплине «Управление проектами»', c, indent=False)
        self.p(f'на тему «{topic}»', c, indent=False)
        for _ in range(8):
            self.blank()
        tab_pos = Cm(TEXT_WIDTH_CM)
        for label, name in ((f'Выполнила: слушатель гр. {group}', student), ('Проверила:', teacher)):
            par = self.p(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            par.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT)
            _font(par.add_run(f'{label}\t{name}'))
            self.blank()
        for _ in range(6):
            self.blank()
        last = self.p(f'Минск {year}', c, indent=False)
        last.add_run().add_break(WD_BREAK.PAGE)

    def contents(self):
        self.p('СОДЕРЖАНИЕ', WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True)
        self.blank()
        par = self.p(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
        _field(par, 'TOC \\o "1-2" \\h \\z \\u', 'Содержание формируется в Word: ПКМ → Обновить поле.')

    # --- заголовки ----------------------------------------------------------
    def section(self, title):
        self.section_no += 1
        self.letter = None
        self.fig = self.tab = self.frm = 0
        h = self.doc.add_paragraph(style='Heading 1')
        h.paragraph_format.page_break_before = True
        _font(h.add_run(f'{self.section_no} {title.upper()}'), bold=True)
        self.blank(keep=True)

    def unnumbered(self, title):
        h = self.doc.add_paragraph(style='Heading 1')
        h.paragraph_format.page_break_before = True
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.first_line_indent = Cm(0)
        _font(h.add_run(title.upper()), bold=True)
        self.blank(keep=True)

    def subsection(self, number, title):
        if len(self.doc.paragraphs) and self.doc.paragraphs[-1].text.strip():
            self.blank(keep=True)
        h = self.doc.add_paragraph(style='Heading 2')
        _font(h.add_run(f'{self.section_no}.{number} {title}'), bold=True)
        self.blank(keep=True)

    def appendix(self, letter, kind, title):
        self.letter = letter
        self.fig = self.tab = 0
        h = self.doc.add_paragraph(style='Heading 1')
        h.paragraph_format.page_break_before = True
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.first_line_indent = Cm(0)
        _font(h.add_run(f'ПРИЛОЖЕНИЕ {letter}'), bold=True)
        h.add_run().add_break()
        _font(h.add_run(f'({kind})'), bold=True)
        h.add_run().add_break()
        _font(h.add_run(title), bold=True)
        self.blank(keep=True)

    # --- текст ------------------------------------------------------------------
    def text(self, text):
        return self.p(text)

    def dash_list(self, items):
        for k, item in enumerate(items):
            end = '.' if k == len(items) - 1 else ';'
            item = item.rstrip(' .;')
            self.p(f'– {item}{end}')

    def num_list(self, items):
        for k, item in enumerate(items, 1):
            self.p(f'{k} {item}')

    def _label(self, kind):
        prefix = self.letter if self.letter else str(self.section_no)
        return f'{prefix}.{kind}'

    def next_fig(self):
        return self._label(self.fig + 1)

    def next_tab(self):
        return self._label(self.tab + 1)

    def figure(self, path, caption, width_cm=16.0):
        self.fig += 1
        self.blank(keep=True)
        par = self.p(align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, keep=True)
        par.add_run().add_picture(path, width=Cm(min(width_cm, TEXT_WIDTH_CM)))
        self.blank(keep=True)
        self.p(f'Рисунок {self._label(self.fig)} – {caption}', WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        self.blank()

    def table(self, caption, headers, rows, widths_cm=None, size=12, header_bold=True):
        self.tab += 1
        if self.doc.paragraphs[-1].text.strip():
            self.blank(keep=True)
        cap = self.p(f'Таблица {self._label(self.tab)} – {caption}', WD_ALIGN_PARAGRAPH.LEFT, indent=False, keep=True)
        t = self.doc.add_table(rows=1, cols=len(headers))
        _borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _repeat_header(t.rows[0])
        for i, h in enumerate(headers):
            cp = t.rows[0].cells[i].paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.first_line_indent = Cm(0)
            _font(cp.add_run(str(h)), size, header_bold)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cp = cells[i].paragraphs[0]
                cp.paragraph_format.first_line_indent = Cm(0)
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _font(cp.add_run(str(val)), size)
        for n, row in enumerate(t.rows):
            _cant_split(row)
            if n <= 2:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        if widths_cm:
            total = sum(widths_cm)
            scale = min(1.0, TEXT_WIDTH_CM / total)
            for row in t.rows:
                for cell, w in zip(row.cells, widths_cm):
                    cell.width = Cm(w * scale)
        self.blank()
        return cap

    def formula(self, expr, explain=None):
        self.frm += 1
        par = self.p(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
        ts = par.paragraph_format.tab_stops
        ts.add_tab_stop(Cm(TEXT_WIDTH_CM / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        _font(par.add_run(f'\t{expr},\t({self._label(self.frm)})'), italic=False)
        if explain:
            self.p(explain, indent=False)
        return self._label(self.frm)

    def listing(self, code_text, size=10):
        for line in code_text.splitlines():
            par = self.p(indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            _font(par.add_run(line if line else ' '), size, name='Courier New')

    # --- источники -------------------------------------------------------------
    def cite(self, key, description):
        for k, (kk, _) in enumerate(self.sources, 1):
            if kk == key:
                return f'[{k}]'
        self.sources.append((key, description))
        return f'[{len(self.sources)}]'

    def bibliography(self):
        self.unnumbered('Список использованных источников')
        for k, (_, desc) in enumerate(self.sources, 1):
            self.p(f'[{k}] {desc}')

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.doc.save(path)
