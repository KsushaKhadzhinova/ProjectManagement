import os
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from template_fill import FONT, set_para, style_run

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, 'ПЗ1_Концепция-проекта', '02_задание-оригиналы', 'Шаблон ПЗ1 - Концепция проекта (шаблон пз1 -УП- 26).docx')
TITLE_LAST = 43


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tbl_pr.append(borders)


class Report:
    def __init__(self, number, topic):
        self.doc = docx.Document(TEMPLATE)
        paras = list(self.doc.paragraphs)
        set_para(paras[16], f'По практическому занятию №{number}')
        set_para(paras[19], f'на тему «{topic}»')
        set_para(paras[28], 'Слушатель гр.60121    \t        К.А. Хаджинова')
        body = self.doc.element.body
        keep = {p._p for p in paras[:TITLE_LAST + 1]}
        for el in list(body.iterchildren()):
            if el.tag.endswith('}sectPr') or el in keep:
                continue
            body.remove(el)
        paras[TITLE_LAST].add_run().add_break(WD_BREAK.PAGE)
        self.figure_no = 0
        self.table_no = 0

    def _para(self, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True):
        p = self.doc.add_paragraph()
        p.alignment = align
        fmt = p.paragraph_format
        fmt.first_line_indent = Cm(1.25) if indent else Cm(0)
        fmt.left_indent = Cm(0)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        return p

    def heading1(self, text, new_page=True):
        p = self._para(indent=True)
        if new_page and len(self.doc.paragraphs) > TITLE_LAST + 2:
            p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(14)
        style_run(p.add_run(text.upper()), 14, bold=True)

    def heading2(self, text):
        p = self._para()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(14)
        style_run(p.add_run(text), 14, bold=True)

    def text(self, text):
        p = self._para()
        style_run(p.add_run(text), 14)

    def bullets(self, items):
        for item in items:
            p = self._para()
            style_run(p.add_run('– ' + item), 14)

    def figure(self, path, caption, width_cm=16.5):
        self.figure_no += 1
        p = self._para(WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width_cm))
        cap = self._para(WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(14)
        style_run(cap.add_run(f'Рисунок {self.figure_no} – {caption}'), 14)
        return self.figure_no

    def next_figure(self):
        return self.figure_no + 1

    def next_table(self):
        return self.table_no + 1

    def table(self, caption, headers, rows, widths_cm=None, size=12):
        self.table_no += 1
        cap = self._para(WD_ALIGN_PARAGRAPH.LEFT, indent=False)
        cap.paragraph_format.space_before = Pt(14)
        cap.paragraph_format.keep_with_next = True
        style_run(cap.add_run(f'Таблица {self.table_no} – {caption}'), 14)
        t = self.doc.add_table(rows=1, cols=len(headers))
        set_borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cp = t.rows[0].cells[i].paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_run(cp.add_run(str(h)), size, bold=True)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cp = cells[i].paragraphs[0]
                cp.paragraph_format.first_line_indent = Cm(0)
                style_run(cp.add_run(str(val)), size)
        if widths_cm:
            for row in t.rows:
                for cell, w in zip(row.cells, widths_cm):
                    cell.width = Cm(w)
        spacer = self._para(indent=False)
        spacer.paragraph_format.space_after = Pt(6)
        return self.table_no

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.doc.save(path)
