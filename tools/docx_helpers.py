# -*- coding: utf-8 -*-
"""Общие помощники для сборки отчётов в стиле БГУИР (по образцам её реальных отчётов)."""
import re
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


LATIN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9_./+#\-]*")


def new_document():
    doc = docx.Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), 'Times New Roman')
    sec = doc.sections[0]
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(1)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    return doc


def add_mixed_para(doc, text, size=14, align=None, space_after=None, italic_all=False):
    """Пишет абзац, автоматически выделяя латинские термины курсивом
    (стиль: 'латиница курсивом даже посреди русского предложения')."""
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(1.25) if align is None else None

    pos = 0
    for m in LATIN_RE.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos:m.start()])
            run.font.size = Pt(size)
            run.font.name = 'Times New Roman'
            run.italic = italic_all
        run = p.add_run(m.group())
        run.font.size = Pt(size)
        run.font.name = 'Times New Roman'
        run.italic = True
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        run.font.size = Pt(size)
        run.font.name = 'Times New Roman'
        run.italic = italic_all
    return p


def add_plain_para(doc, text, bold=False, italic=False, size=14, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    if align:
        p.alignment = align
    return p


def _style_heading_run(p, text, size=14):
    for run in p.runs:
        run.text = ""
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0, 0, 0)
    return run


def add_heading1(doc, text):
    """Настоящий Heading 1 (нужен для автосборки TOC), но оформлен под стиль отчёта."""
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    _style_heading_run(p, text.upper())
    return p


def add_heading2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    _style_heading_run(p, text)
    return p


def add_figure(doc, image_path, caption, width_cm=15):
    doc.add_picture(image_path, width=Cm(width_cm))
    last_p = doc.paragraphs[-1]
    last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    cap.paragraph_format.space_after = Pt(12)


def add_table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(12)
            run.font.name = 'Times New Roman'
    doc.add_paragraph("")
    return t


def add_toc_field(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char_begin = OxmlElement('w:fldChar')
    fld_char_begin.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_char_separate = OxmlElement('w:fldChar')
    fld_char_separate.set(qn('w:fldCharType'), 'separate')
    fld_char_text = OxmlElement('w:t')
    fld_char_text.text = "Нажми правой кнопкой -> Обновить поле, чтобы собрать оглавление"
    fld_char_end = OxmlElement('w:fldChar')
    fld_char_end.set(qn('w:fldCharType'), 'end')

    r_element = run._r
    r_element.append(fld_char_begin)
    r_element.append(instr_text)
    r_element.append(fld_char_separate)
    r_element.append(fld_char_text)
    r_element.append(fld_char_end)
    return p


def title_page(doc, topic_number, topic_name, on_topic=None):
    for line in [
        "Министерство образования Республики Беларусь",
        "Учреждение образования",
        "«Белорусский государственный университет информатики",
        "и радиоэлектроники»",
        "Обособленное подразделение",
        "«Институт информационных технологий БГУИР»",
        "Факультет повышения квалификации и переподготовки",
        "Кафедра микропроцессорных систем и сетей",
    ]:
        add_plain_para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER)

    for _ in range(4):
        doc.add_paragraph("")

    add_plain_para(doc, "ОТЧЁТ", bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_plain_para(doc, f"по практическому занятию {topic_number}", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_plain_para(doc, "по дисциплине «Управление проектами»", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_plain_para(doc, f"на тему «{topic_name}»", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    if on_topic:
        add_plain_para(doc, on_topic, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    for _ in range(4):
        doc.add_paragraph("")

    add_plain_para(doc, "Слушатель гр. 60121" + " " * 55 + "Хаджинова К. А.")
    add_plain_para(doc, "Руководитель" + " " * 65 + "И. В. Кашникова")

    for _ in range(6):
        doc.add_paragraph("")
    add_plain_para(doc, "МИНСК 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def toc_page(doc):
    add_plain_para(doc, "СОДЕРЖАНИЕ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph("")
    add_toc_field(doc)
    doc.add_page_break()
