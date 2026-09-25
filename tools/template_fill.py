import copy
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = 'Times New Roman'


def style_run(run, size=14, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def set_para(p, text, size=None, bold=None, italic=None):
    runs = p.runs
    if runs:
        first = runs[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        first.text = text
        if size:
            first.font.size = Pt(size)
        if bold is not None:
            first.bold = bold
        if italic is not None:
            first.italic = italic
        first.font.name = FONT
    else:
        style_run(p.add_run(text), size or 14, bold, italic)
    return p


def set_cell(cell, text, size=12, bold=False):
    lines = text if isinstance(text, list) else [text]
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    style_run(p.add_run(lines[0]), size, bold)
    for line in lines[1:]:
        np_ = cell.add_paragraph()
        np_.paragraph_format.first_line_indent = Cm(0)
        style_run(np_.add_run(line), size, bold)


def remove_para(p):
    p._p.getparent().remove(p._p)


def insert_after(anchor_el, doc, text='', size=14, bold=False, italic=False, align=None, indent_cm=None):
    new_p = copy.deepcopy(doc.paragraphs[0]._p)
    for child in list(new_p):
        if not child.tag.endswith('}pPr'):
            new_p.remove(child)
    anchor_el.addnext(new_p)
    p = docx.text.paragraph.Paragraph(new_p, doc.paragraphs[0]._parent)
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent_cm is not None:
        p.paragraph_format.left_indent = Cm(indent_cm)
        p.paragraph_format.first_line_indent = Cm(0)
    if text:
        style_run(p.add_run(text), size, bold, italic)
    return p


def insert_picture_after(anchor_el, doc, path, width_cm, caption):
    pic_p = insert_after(anchor_el, doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    pic_p.paragraph_format.first_line_indent = Cm(0)
    pic_p.add_run().add_picture(path, width=Cm(width_cm))
    cap = insert_after(pic_p._p, doc, caption, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.first_line_indent = Cm(0)
    return cap


def fill_table_rows(table, rows, size=12, start=1):
    while len(table.rows) - start < len(rows):
        new_tr = copy.deepcopy(table.rows[-1]._tr)
        table._tbl.append(new_tr)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            set_cell(table.rows[start + i].cells[j], val, size)
